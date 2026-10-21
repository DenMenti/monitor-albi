"""
logica.py — funzioni "pure" (senza browser): riconoscimento dei link all'albo,
lettura delle date, filtro per parole chiave, nomi dei file, Excel.
"""
import datetime as dt
import re
import unicodedata
from urllib.parse import urlparse

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

# ---------------------------------------------------------------- testo

def norm(s):
    """minuscolo, senza accenti, spazi compattati"""
    s = unicodedata.normalize("NFKD", s or "")
    s = "".join(ch for ch in s if not unicodedata.combining(ch))
    s = s.replace("’", "'")
    return re.sub(r"\s+", " ", s).strip().lower()


def slug(s):
    s = norm(s).replace("'", "-")
    s = re.sub(r"[^a-z0-9]+", "-", s)
    return s.strip("-")


def chiave(c):
    """Identifica un comune in modo univoco: 'Agordo (BL)'."""
    return f"{c['nome']} ({c['sigla']})" if c.get("sigla") else c["nome"]


def normalizza_url(u):
    u = (u or "").strip().rstrip(".,;")
    if not u:
        return ""
    if not re.match(r"^https?://", u, re.I):
        u = "https://" + u
    return u


def href_utile(href):
    if not href:
        return False
    h = href.strip().lower()
    return not (h.startswith("javascript:") or h in ("#", "#!") or h.startswith("mailto:")
                or h.startswith("tel:"))


# ----------------------------------------------------------- date

DATA_RE = re.compile(r"\b(\d{1,2})[/.\-](\d{1,2})[/.\-](\d{4}|\d{2})\b")


def tutte_le_date(s):
    out = []
    for g, m, a in DATA_RE.findall(s or ""):
        a = int(a)
        if a < 100:
            a += 2000
        try:
            out.append(dt.date(a, int(m), int(g)))
        except ValueError:
            pass
    return out


def prima_data(s):
    d = tutte_le_date(s)
    return d[0] if d else None


# --------------------------------------------- link all'albo (fase 1)

ALBI_ESCLUSI = ["fornitor", "associazion", "scrutator", "president", "albo d'oro",
                "volontar", "professionist", "avvocat", "giudici popolari",
                "operatori economici", "imprese", "beneficiari", "albo comunale delle"]


def punteggio_link_albo(c):
    testo = norm(c.get("text", "") + " " + c.get("title", ""))
    href = (c.get("href", "") or "").lower()
    tutto = testo + " " + href
    if "albo" not in tutto and "mc_p_ricerca" not in href:
        return None
    if any(x in tutto for x in ALBI_ESCLUSI):
        return None
    if "matri" in href:
        return None
    s = 0
    if "albo pretorio" in testo:
        s += 10
    if re.search(r"albo\s*on\s*-?\s*line", testo):
        s += 8
    if "mc_p_ricerca" in href:
        s += 7
    if "albo" in testo:
        s += 4
    if "albo" in href:
        s += 2
    if e_archivio(c):
        s -= 6
    return s


def e_archivio(c):
    t = norm(c.get("text", "") + " " + c.get("title", "") + " " + c.get("href", ""))
    return bool(re.search(r"archivio|storic|fino al|precedent", t))


def indovina_piattaforma(urls, info):
    j = " ".join(u for u in urls if u).lower()
    titolo = (info.get("title") or "").lower()
    gen = (info.get("generator") or "").lower()
    crediti = (info.get("credits") or "").lower()
    if "mc_p_ricerca" in j or "halley" in j or "halley" in crediti:
        return "Halley"
    if "myportal" in titolo:
        return "myPortal"
    if "trasparenza-valutazione-merito" in j or "gazzettaamministrativa" in j:
        return "Gazzetta Amministrativa"
    if "municipium" in j:
        return "Municipium"
    if "maggioli" in j or "maggioli" in crediti:
        return "Maggioli"
    if "wordpress" in gen:
        return "WordPress"
    return ""


# ------------------------------------------- righe dell'albo (fase 2)

def _trova_col(headers, *chiavi, escludi=()):
    for i, h in enumerate(headers):
        n = norm(h)
        if any(k in n for k in chiavi) and not any(e in n for e in escludi):
            return i
    return None


def mappa_colonne(headers):
    """Indici delle colonne utili a partire dalle intestazioni della tabella."""
    if not headers:
        return {}
    m = {}
    m["data"] = (_trova_col(headers, "inizio pubbl", "data pubbl", "pubblicato dal", "pubblicazione dal",
                            "periodo", "dal")
                 or _trova_col(headers, "pubblicazion", escludi=("fine", "al ", "scadenza"))
                 or _trova_col(headers, "data", escludi=("fine", "scadenza")))
    m["oggetto"] = _trova_col(headers, "oggetto", "descrizione", "titolo")
    m["tipo"] = _trova_col(headers, "tipo", "tipologia", "categoria", "atto", escludi=("data", "numero", "n."))
    m["numero"] = _trova_col(headers, "numero", "n. reg", "registro", "n.", "num")
    return {k: v for k, v in m.items() if v is not None}


def riga_in_atto(r):
    """Converte una riga grezza letta dalla pagina in un dizionario 'atto'."""
    cells = r.get("cells") or []
    headers = r.get("headers") or []
    testo = r.get("text", "")
    # le intestazioni a volte hanno una colonna in più/meno: allineo solo se coincidono
    m = mappa_colonne(headers) if len(headers) == len(cells) else {}

    def cella(k):
        i = m.get(k)
        return cells[i] if i is not None and i < len(cells) else ""

    data = prima_data(cella("data")) if m.get("data") is not None else None
    if data is None:
        date = tutte_le_date(testo)
        if len(date) >= 3:
            data = date[1]       # spesso: data atto, inizio, fine pubblicazione
        elif date:
            data = date[0]       # inizio (e fine) pubblicazione
    oggetto = cella("oggetto")
    if not oggetto:
        oggetto = max(cells, key=len) if cells else testo
    return {
        "data": data,
        "oggetto": re.sub(r"\s+", " ", oggetto).strip(),
        "tipo": cella("tipo"),
        "numero": cella("numero"),
        "testo": testo,
        "links": r.get("links") or [],
    }


def valuta_atto(atto, parole, esclusioni, solo_determine):
    """Ritorna la lista di parole chiave trovate, oppure None se l'atto va scartato."""
    t = norm(atto["testo"] + " " + atto["oggetto"] + " " + atto["tipo"])
    if solo_determine and "determin" not in t:
        return None
    if any(norm(e) in t for e in esclusioni if e.strip()):
        return None
    parole = [p for p in parole if p.strip()]
    if not parole:
        return ["(nessuna parola chiave: tutte le determine)"]
    trovate = [p for p in parole if norm(p) in t]
    return trovate or None


ESTENSIONI_FILE = (".pdf", ".p7m", ".doc", ".docx", ".odt", ".rtf", ".zip", ".xls", ".xlsx", ".ods",
                   ".jpg", ".jpeg", ".png", ".tif", ".tiff", ".xml")


def sembra_allegato(link):
    href = (link.get("href") or "")
    if not href_utile(href):
        return False
    h = href.lower()
    path = urlparse(h).path
    if path.endswith(ESTENSIONI_FILE):
        return True
    t = norm(link.get("text", ""))
    if re.search(r"allegat|download|scarica|attachment|mc_attachment|visualizza documento|documento principale"
                 r"|getfile|file\.php|downloadfile|apri file", h + " " + t):
        return True
    return False


def link_dettaglio(atto):
    """Il link migliore per aprire il dettaglio dell'atto (non un allegato)."""
    candidati = [l for l in atto["links"] if href_utile(l.get("href")) and not sembra_allegato(l)]
    if not candidati:
        return ""
    for l in candidati:
        if re.search(r"dettagli|visualizza|apri|mostra|leggi", norm(l.get("text", "")) + " " + l["href"].lower()):
            return l["href"]
    return candidati[0]["href"]


EST_DA_MIME = {"application/pdf": ".pdf", "application/pkcs7-mime": ".p7m", "application/x-pkcs7-mime": ".p7m",
               "application/msword": ".doc", "application/zip": ".zip",
               "application/vnd.openxmlformats-officedocument.wordprocessingml.document": ".docx",
               "application/vnd.oasis.opendocument.text": ".odt", "image/jpeg": ".jpg", "image/png": ".png"}


def estensione(url, content_type="", content_disposition=""):
    m = re.search(r'filename\*?=(?:UTF-8\'\')?"?([^";]+)', content_disposition or "", re.I)
    for nome in ([m.group(1)] if m else []) + [urlparse(url).path]:
        nome = nome.lower()
        if nome.endswith(".pdf.p7m"):
            return ".pdf.p7m"
        for e in ESTENSIONI_FILE:
            if nome.endswith(e):
                return e
    return EST_DA_MIME.get((content_type or "").split(";")[0].strip().lower(), ".pdf")


def nome_file(comune, data, n, est):
    d = data.isoformat() if data else "senza-data"
    return f"{slug(comune)}_{d}_{n:02d}{est}"


# ---------------------------------------------------------------- Excel

BLU = "2E5E4E"
GIALLO = PatternFill("solid", fgColor="FFF2CC")
F = Font(name="Arial", size=10)
LINK = Font(name="Arial", size=10, color="0563C1", underline="single")


def _intesta(ws, colonne, larghezze):
    ws.append(colonne)
    for c in ws[1]:
        c.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=BLU)
        c.alignment = Alignment(vertical="center")
    for i, w in enumerate(larghezze):
        ws.column_dimensions[chr(65 + i)].width = w
    ws.freeze_panes = "A2"


def _riga(ws, valori, link_cols=(), evidenzia=False):
    ws.append(valori)
    r = ws.max_row
    for c in ws[r]:
        c.font = F
        c.alignment = Alignment(vertical="top", wrap_text=True)
        if evidenzia:
            c.fill = GIALLO
    for col, target in link_cols:
        if target:
            cella = ws.cell(row=r, column=col)
            cella.hyperlink = target
            cella.font = LINK


def excel_albi(comuni, percorso):
    wb = Workbook()
    ws = wb.active
    ws.title = "Albi pretori"
    _intesta(ws, ["Comune", "Provincia", "Regione", "Sito istituzionale", "Link Albo Pretorio", "Altri link albo",
                  "Piattaforma", "Note"], [26, 10, 16, 38, 60, 50, 18, 45])
    for c in sorted(comuni, key=lambda c: (c.get("regione", ""), c.get("sigla", ""), c["nome"])):
        albo = c.get("albo", "")
        _riga(ws, [c["nome"], c.get("sigla", ""), c.get("regione", ""), c.get("sito", ""), albo,
                   "\n".join(c.get("altri", [])), c.get("piattaforma", ""), c.get("note", "")],
              link_cols=[(4, c.get("sito")), (5, albo)],
              evidenzia=not albo or "verifica" in c.get("note", ""))
    ws.auto_filter.ref = ws.dimensions
    wb.save(percorso)


def excel_monitoraggio(risultati, riepilogo, percorso):
    wb = Workbook()
    ws = wb.active
    ws.title = "Determine trovate"
    _intesta(ws, ["Comune", "Provincia", "Data pubblicazione", "Tipo", "Numero", "Oggetto", "Parole chiave",
                  "Link atto", "Allegati"], [22, 10, 14, 18, 12, 70, 24, 40, 40])
    for r in risultati:
        primo = r["allegati"][0] if r["allegati"] else ""
        _riga(ws, [r["comune"], r.get("sigla", ""), r["data"].strftime("%d/%m/%Y") if r["data"] else "", r["tipo"],
                   r["numero"], r["oggetto"], ", ".join(r["parole"]), r["link"], "\n".join(r["allegati"])],
              link_cols=[(8, r["link"]), (9, primo)],
              evidenzia=not r["allegati"])
    if ws.max_row > 1:
        ws.auto_filter.ref = ws.dimensions

    ws2 = wb.create_sheet("Riepilogo comuni")
    _intesta(ws2, ["Comune", "Provincia", "Link albo", "Controllato dal", "Pagine lette", "Atti letti", "Trovati",
                   "Esito"], [22, 10, 50, 15, 12, 10, 10, 50])
    for r in riepilogo:
        ok = r["esito"].startswith("OK")
        _riga(ws2, [r["comune"], r.get("sigla", ""), r["albo"], r["dal"].strftime("%d/%m/%Y") if r["dal"] else "",
                    r["pagine"], r["letti"], r["trovati"], r["esito"]],
              link_cols=[(3, r["albo"])], evidenzia=not ok)
    ws2.append([])
    ws2.append(["Righe in giallo: albo non letto o errore, da controllare a mano."])
    ws2.cell(row=ws2.max_row, column=1).font = Font(name="Arial", size=9, italic=True)
    wb.save(percorso)
